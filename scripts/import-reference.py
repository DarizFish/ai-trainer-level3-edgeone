from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

from docx import Document

from notebook_solutions import (
    DEFAULT_AUDIT_PATH as DEFAULT_NOTEBOOK_AUDIT_PATH,
    DEFAULT_SNAPSHOT_PATH as DEFAULT_NOTEBOOK_SNAPSHOT_PATH,
    apply_notebook_solutions,
    content_metadata as notebook_solution_metadata,
    load_snapshot as load_notebook_solution_snapshot,
)
from practical_spreadsheet_answers import (
    DEFAULT_SNAPSHOT_PATH as DEFAULT_PRACTICAL_SPREADSHEET_PATH,
    apply_spreadsheet_answers,
    content_metadata as practical_spreadsheet_metadata,
    load_snapshot as load_practical_spreadsheet_snapshot,
)
from practical_template_answers import (
    DEFAULT_SNAPSHOT_PATH as DEFAULT_PRACTICAL_TEMPLATE_PATH,
    apply_template_answers,
    content_metadata as practical_template_metadata,
    load_snapshot as load_practical_template_snapshot,
)
from practical_rubric_answers import (
    DEFAULT_SNAPSHOT_PATH as DEFAULT_PRACTICAL_RUBRIC_PATH,
    apply_rubric_answers,
    content_metadata as practical_rubric_metadata,
    load_snapshot as load_practical_rubric_snapshot,
)
from theory_reference import (
    DEFAULT_AUDIT_PATH,
    DEFAULT_EDITORIAL_REVIEW_PATH,
    DEFAULT_REFERENCE_PATH,
    apply_editorial_reviews,
    apply_reference_answers,
    content_reference_metadata,
    load_editorial_review,
    load_reference,
    write_json,
)


VERSION = "2026-07-19-v7"
TYPE_LABELS = {"judge": "判断题", "single": "单选题", "multi": "多选题"}
CATEGORY_RANGES = [
    (1, 45, "基本要求", "职业道德与基础知识"),
    (46, 105, "业务分析", "业务流程与模块优化"),
    (106, 195, "智能训练", "数据处理与算法测试"),
    (196, 285, "智能系统设计", "系统优化与人机交互"),
    (286, 300, "培训与指导", "培训实施与问题指导"),
]
GROUPS = ["1.1", "1.2", "2.1", "2.2", "3.1", "3.2", "4.1", "4.2"]
GROUP_META = {
    "1.1": ("业务分析", "业务流程设计", 25, 30),
    "1.2": ("业务分析", "业务模块效果优化", 25, 30),
    "2.1": ("智能训练", "数据处理规范制定", 15, 20),
    "2.2": ("智能训练", "算法测试", 20, 20),
    "3.1": ("智能系统设计", "智能系统监控和优化", 15, 20),
    "3.2": ("智能系统设计", "人机交互流程设计", 20, 20),
    "4.1": ("培训与指导", "培训", 5, 10),
    "4.2": ("培训与指导", "指导", 5, 10),
}
OFFICIAL_PRACTICAL_CODES = {"1.1.1", "2.1.1", "2.2.1", "3.1.1", "3.2.1", "4.2.1"}


def clean_official_practical_answer(value: str) -> str:
    """Keep one official answer while removing Word page furniture and controls.

    The legacy Word document places the next task sheet after a form-feed.  A
    task name appears later in that sheet, so stopping at the next task name
    alone leaks the next sheet header into the current answer.
    """
    boundary = re.search(
        r"\f(?=\s*人工智能训练师（三级）操作技能考核\s*试题单\b)",
        value,
    )
    if boundary:
        value = value[: boundary.start()]
    value = value.replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[\x07\x0b]", "\n", value)
    value = re.sub(r"[\x00-\x06\x08\x0c\x0e-\x1f\x7f]", "", value)
    return re.sub(r"\n{3,}", "\n\n", value).strip()


def compact(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    return re.sub(r"\s+", "", value)


def strip_number(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).strip()
    value = re.sub(r"^[（(]\s*[）)]\s*", "", value)
    value = re.sub(r"^.*?\d+\s*[.、]\s*", "", value, count=1)
    return value.strip()


def normalize_stem(value: str) -> str:
    return compact(strip_number(value)).replace("(", "").replace(")", "")


def option_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).strip()
    return re.sub(r"^[（(]\s*[A-E]\s*[）)]\s*", "", value).strip()


def category_for(index: int) -> tuple[str, str]:
    for start, end, category, section in CATEGORY_RANGES:
        if start <= index <= end:
            return category, section
    raise ValueError(index)


def hash_choice(text: str, count: int) -> int:
    digest = hashlib.sha256(text.encode("utf-8")).digest()
    return int.from_bytes(digest[:2], "big") % count


FALSE_MARKERS = (
    "无需", "不需要", "完全", "只能", "仅仅", "唯一", "所有", "一定", "不存在",
    "无关", "不能", "不涉及", "没有必要", "随意", "直接删除", "100%", "始终",
    "任何情况下", "可以不经", "不必", "固定不变", "完全替代", "唯一方法",
)
BAD_OPTION_MARKERS = (
    "随意", "忽略", "无需", "不需要", "完全", "唯一", "仅", "删除所有", "不考虑",
    "无法", "不能", "越大越好", "只关注", "任意", "省略", "夸大", "泄露",
)


def derived_answer(kind: str, stem: str, options: list[str]) -> tuple[list[str], str, str]:
    if kind == "judge":
        is_false = any(marker in stem for marker in FALSE_MARKERS)
        answer = ["×" if is_false else "√"]
        reason = "题干含绝对化或忽略必要条件的表述。" if is_false else "题干与该知识点的一般定义和规范相符。"
        return answer, reason, "low"

    negative = any(word in stem for word in ("不属于", "不是", "错误", "无效", "不能", "不包括"))
    letters = [chr(65 + i) for i in range(len(options))]
    if kind == "single":
        bad = [i for i, text in enumerate(options) if any(marker in text for marker in BAD_OPTION_MARKERS)]
        if negative and bad:
            index = bad[0]
        elif not negative and len(bad) == len(options) - 1:
            index = next(i for i in range(len(options)) if i not in bad)
        else:
            stem_chars = set(compact(stem))
            scores = [len(stem_chars & set(compact(text))) for text in options]
            target = min(scores) if negative else max(scores)
            candidates = [i for i, score in enumerate(scores) if score == target]
            index = candidates[hash_choice(stem, len(candidates))]
        return [letters[index]], "依据题干关键词、选项语义及常规知识定义整理。", "low"

    selected = []
    for i, text in enumerate(options):
        is_bad = any(marker in text for marker in BAD_OPTION_MARKERS)
        if (negative and is_bad) or (not negative and not is_bad):
            selected.append(letters[i])
    if len(selected) < 2:
        start = hash_choice(stem, max(1, len(options) - 1))
        selected = letters[start : start + 2]
    return selected, "依据各选项是否同时满足题干条件整理，多选须完全匹配。", "low"


def parse_theory(source: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    review_path = next(source.glob("第3部分*.docx"))
    mock_path = next(source.glob("第5部分*.docx"))
    review_paragraphs = [p.text.strip() for p in Document(review_path).paragraphs if p.text.strip()]
    mock_paragraphs = [p.text.strip() for p in Document(mock_path).paragraphs if p.text.strip()]

    raw: dict[str, list[tuple[str, list[str]]]] = {
        "judge": [(review_paragraphs[4 + i], []) for i in range(300)],
        "single": [
            (review_paragraphs[305 + i * 5], [option_text(x) for x in review_paragraphs[306 + i * 5 : 310 + i * 5]])
            for i in range(300)
        ],
        "multi": [
            (review_paragraphs[1806 + i * 6], [option_text(x) for x in review_paragraphs[1807 + i * 6 : 1812 + i * 6]])
            for i in range(300)
        ],
    }

    judge_answers: list[str] = []
    for line in mock_paragraphs[813:817]:
        judge_answers.extend(re.findall(r"[√×]", line))
    single_answers: list[str] = []
    for line in mock_paragraphs[817:831]:
        match = re.search(r"([A-D]{10})\s*$", compact(line))
        if match:
            single_answers.extend(list(match.group(1)))
    multi_answers = []
    for line in mock_paragraphs[832:842]:
        match = re.search(r"\d+[)）]\s*([A-E]+)", line)
        if match:
            multi_answers.append(list(match.group(1)))

    mock_raw = {
        "judge": [(mock_paragraphs[9 + i], [judge_answers[i]]) for i in range(40)],
        "single": [(mock_paragraphs[50 + i * 5], [single_answers[i]]) for i in range(140)],
        "multi": [(mock_paragraphs[751 + i * 6], multi_answers[i]) for i in range(10)],
    }
    official_by_stem = {
        kind: {normalize_stem(stem): answer for stem, answer in values}
        for kind, values in mock_raw.items()
    }

    questions: list[dict[str, Any]] = []
    ids_by_stem: dict[str, dict[str, str]] = {kind: {} for kind in raw}
    for kind, values in raw.items():
        for index, (stem, options) in enumerate(values, 1):
            normalized = normalize_stem(stem)
            official = official_by_stem[kind].get(normalized)
            if official:
                answer, explanation, confidence = official, "答案来自指导手册理论知识模拟试卷。", "high"
                source_type = "official"
                answer_basis = "manual_official"
            else:
                answer, explanation, confidence = derived_answer(kind, strip_number(stem), options)
                source_type = "derived"
                answer_basis = "heuristic"
            category, section = category_for(index)
            question_id = f"{kind}-{index:03d}"
            ids_by_stem[kind][normalized] = question_id
            questions.append({
                "id": question_id,
                "index": index,
                "type": kind,
                "typeLabel": TYPE_LABELS[kind],
                "category": category,
                "section": section,
                "stem": strip_number(stem),
                "options": options,
                "answer": answer,
                "answerSource": source_type,
                "answerBasis": answer_basis,
                "confidence": confidence,
                "explanation": explanation,
                "points": 1 if kind == "multi" else 0.5,
            })

    official_ids = {
        kind: [ids_by_stem[kind][normalize_stem(stem)] for stem, _ in values]
        for kind, values in mock_raw.items()
    }
    mock = {
        "id": "official-theory-1",
        "title": "理论知识官方模拟试卷",
        "durationMinutes": 90,
        "passingScore": 60,
        "questionIds": official_ids["judge"] + official_ids["single"] + official_ids["multi"],
        "distribution": {"judge": 40, "single": 140, "multi": 10},
        "maxScore": 100,
    }
    return questions, mock


def clean_lines(value: str) -> list[str]:
    return [line.strip(" \t\x07\x0b\r") for line in value.split("\n") if line.strip(" \t\x07\x0b\r")]


def extract_name(block: str) -> str | None:
    match = re.search(r"试题名称：([^\r\n\f]+)", block)
    return match.group(1).strip() if match else None


def parse_rubric(block: str, total: int) -> list[dict[str, Any]]:
    lines = clean_lines(block)
    items: list[dict[str, Any]] = []
    for i, line in enumerate(lines):
        if not re.fullmatch(r"M\d+", line):
            continue
        points = None
        description = None
        for candidate in lines[i + 1 : i + 8]:
            if points is None and re.fullmatch(r"\d+(?:\.\d+)?", candidate):
                points = float(candidate)
                continue
            if points is not None and not re.fullmatch(r"\d+(?:\.\d+)?", candidate) and candidate not in {"根据数据", "规定或", "标称值", "结果或", "实际值", "得分"}:
                description = candidate
                break
        if points is not None and description:
            items.append({"id": line, "points": points, "description": description, "mode": "auto" if any(x in description for x in ("代码", "数据集", "模型", "运行", "保存")) else "self"})
    if not items:
        return [{"id": "M1", "points": total, "description": "按题目要求与评分标准完成任务。", "mode": "self"}]
    delta = total - sum(item["points"] for item in items)
    if abs(delta) > 0.01:
        items.append({"id": "M-final", "points": delta, "description": "其余质量指标与提交规范。", "mode": "self"})
    return items


def task_sheet_section(block: str, start: str, end: str | None) -> str:
    """Extract one non-scoring section from the official task sheet.

    `description` remains the fingerprinted work-task body. The surrounding
    task-sheet sections are kept separately for the candidate-facing full
    question view, without changing audited answer-reference bindings.
    """
    value = block.split(start, 1)[1] if start in block else ""
    if end and end in value:
        value = value.split(end, 1)[0]
    value = re.sub(r"\x0b", "\n", value)
    return "\n".join(clean_lines(value))


def notebook_payload(materials: Path, code: str) -> tuple[list[dict[str, str]], list[str]]:
    folder = materials / f"{code}-素材"
    if not folder.exists():
        return [], []
    assets = [p.name for p in sorted(folder.iterdir()) if p.is_file() and p.suffix.lower() != ".ipynb"]
    notebook = next(folder.glob("*.ipynb"), None)
    if notebook is None:
        return [], assets
    data = json.loads(notebook.read_text(encoding="utf-8"))
    cells = []
    for index, cell in enumerate(data.get("cells", []), 1):
        if cell.get("cell_type") == "code":
            cells.append({"id": f"cell-{index}", "source": "".join(cell.get("source", []))})
    return cells, assets


def parse_practical(source: Path, extracted: Path) -> list[dict[str, Any]]:
    review = (extracted / "part4.txt").read_text(encoding="utf-8")
    mock = (extracted / "part6.txt").read_text(encoding="utf-8")
    materials = source / "人工智能训练师_3级_sucai"

    question_blocks: dict[str, str] = {}
    score_blocks: dict[str, str] = {}
    names: list[str] = []
    for block in review.split("\f"):
        name = extract_name(block)
        if not name:
            continue
        if "试题评分表" in block or "测量分评分表" in block:
            score_blocks.setdefault(name, block)
        elif "试题单" in block:
            question_blocks.setdefault(name, block)
            if name not in names:
                names.append(name)

    all_occurrences = list(re.finditer(r"试题名称：([^\r\n\f]+)", review))
    for occurrence in all_occurrences:
        name = occurrence.group(1).strip()
        if name not in names:
            names.append(name)
    names.sort(key=lambda value: next(m.start() for m in all_occurrences if m.group(1).strip() == value))

    for name in names:
        positions = [m.start() for m in all_occurrences if m.group(1).strip() == name]
        if positions and name not in question_blocks:
            end = positions[1] if len(positions) > 1 else min(len(review), positions[0] + 12000)
            question_blocks[name] = review[positions[0] : end]
        if len(positions) > 1 and name not in score_blocks:
            later = [m.start() for m in all_occurrences if m.start() > positions[1]]
            end = min(later) if later else len(review)
            score_blocks[name] = review[positions[1] : end]

    # Each of the six official answers follows the second occurrence of its task
    # name (the scoring sheet). Map an answer marker to the nearest preceding
    # task-name marker, then stop at the next task. Searching forward from the
    # table of contents would incorrectly map every task to the first answer.
    official_answers: dict[str, str] = {}
    answer_label = "\u53c2\u8003\u7b54\u6848"
    name_pattern = re.compile(r"\u8bd5\u9898\u540d\u79f0\uff1a?([^\r\n\f]+)")
    mock_name_matches = list(name_pattern.finditer(mock))
    for answer_match in re.finditer(answer_label, mock):
        preceding = [match for match in mock_name_matches if match.start() < answer_match.start()]
        if not preceding:
            continue
        name = preceding[-1].group(1).strip()
        following = [match.start() for match in mock_name_matches if match.start() > answer_match.end()]
        end = min(following) if following else len(mock)
        answer_text = clean_official_practical_answer(mock[answer_match.end() : end])
        if not answer_text:
            raise ValueError(f"Official practical answer is empty: {name}")
        if name in official_answers:
            raise ValueError(f"Duplicate official practical answer: {name}")
        official_answers[name] = answer_text

    tasks = []
    for position, name in enumerate(names):
        group = GROUPS[position // 5]
        code = f"{group}.{position % 5 + 1}"
        category, unit, total, minutes = GROUP_META[group]
        question_block = question_blocks.get(name, "")
        score_block = score_blocks.get(name, "")
        rubrics = parse_rubric(score_block, total)
        cells, assets = notebook_payload(materials, code)
        body = question_block
        if "2.工作任务" in body:
            body = body.split("2.工作任务", 1)[1]
        body = re.split(r"\n\s*3\.技能要求", body, maxsplit=1)[0]
        body = re.sub(r"\x0b", "\n", body)
        body = "\n".join(clean_lines(body))
        task_sheet = {
            "equipmentRequirements": task_sheet_section(question_block, "1.场地设备要求", "2.工作任务"),
            "skillRequirements": task_sheet_section(question_block, "3.技能要求", "4.质量指标"),
            "qualityIndicators": task_sheet_section(question_block, "4.质量指标", None),
        }
        if not all(task_sheet.values()):
            raise ValueError(f"Incomplete official task sheet: {code} {name}")
        if code in OFFICIAL_PRACTICAL_CODES:
            reference = official_answers.get(name)
            if not reference:
                raise ValueError(f"Missing official practical answer for {code}: {name}")
            answer_source = "official"
            answer_basis = "manual_official"
        else:
            reference = "建议答案应完整覆盖以下评分点：\n" + "\n".join(
                f"{item['id']}（{item['points']:g}分）：{item['description']}" for item in rubrics
            )
            if cells:
                reference += "\n\nPython 部分：补全模板中的 ___，逐项运行并确保数据、模型或输出文件满足上述评分点。"
            answer_source = "derived"
            answer_basis = "completed_reference_notebook" if cells else "official_rubric"
        tasks.append({
            "code": code,
            "title": name,
            "category": category,
            "unit": unit,
            "group": group,
            "durationMinutes": minutes,
            "maxScore": total,
            "description": body,
            "officialTaskSheet": task_sheet,
            "rubric": rubrics,
            "answerSource": answer_source,
            "answerBasis": answer_basis,
            "referenceAnswer": reference,
            "materials": [],
            "hasNotebook": bool(cells),
            "notebook": {"cells": cells, "assets": assets, "browserAdapted": group == "3.2"} if cells else None,
        })
    return tasks


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--extracted", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--theory-reference", type=Path, default=DEFAULT_REFERENCE_PATH)
    parser.add_argument("--theory-reference-audit", type=Path, default=DEFAULT_AUDIT_PATH)
    parser.add_argument("--theory-editorial-review", type=Path, default=DEFAULT_EDITORIAL_REVIEW_PATH)
    parser.add_argument("--notebook-solutions", type=Path, default=DEFAULT_NOTEBOOK_SNAPSHOT_PATH)
    parser.add_argument("--notebook-solution-audit", type=Path, default=DEFAULT_NOTEBOOK_AUDIT_PATH)
    parser.add_argument("--practical-spreadsheet-answers", type=Path, default=DEFAULT_PRACTICAL_SPREADSHEET_PATH)
    parser.add_argument("--practical-template-answers", type=Path, default=DEFAULT_PRACTICAL_TEMPLATE_PATH)
    parser.add_argument("--practical-rubric-answers", type=Path, default=DEFAULT_PRACTICAL_RUBRIC_PATH)
    args = parser.parse_args()

    theory, official_mock = parse_theory(args.source)
    reference_snapshot = load_reference(args.theory_reference)
    reference_audit = apply_reference_answers(theory, reference_snapshot)
    editorial_review = load_editorial_review(args.theory_editorial_review)
    editorial_audit = apply_editorial_reviews(theory, editorial_review)
    reference_audit["editorialReview"] = editorial_audit
    reference_audit["summary"]["appliedEditorialReviews"] = editorial_audit["appliedReviews"]
    reference_audit["summary"]["remainingHeuristicAnswers"] = sum(question["answerBasis"] == "heuristic" for question in theory)
    practical = parse_practical(args.source, args.extracted)
    practical_spreadsheet_snapshot = load_practical_spreadsheet_snapshot(args.practical_spreadsheet_answers)
    practical_spreadsheet_audit = apply_spreadsheet_answers(practical, args.source, practical_spreadsheet_snapshot)
    practical_template_snapshot = load_practical_template_snapshot(args.practical_template_answers)
    practical_template_audit = apply_template_answers(practical, args.source, practical_template_snapshot)
    practical_rubric_snapshot = load_practical_rubric_snapshot(args.practical_rubric_answers)
    practical_rubric_audit = apply_rubric_answers(practical, practical_rubric_snapshot)
    notebook_snapshot = load_notebook_solution_snapshot(args.notebook_solutions)
    notebook_audit = apply_notebook_solutions(practical, notebook_snapshot)
    payload = {
        "version": VERSION,
        "generatedFrom": "人工智能训练师（三级）评价指导手册",
        "theoryReference": content_reference_metadata(reference_snapshot, reference_audit),
        "notebookSolutions": notebook_solution_metadata(notebook_snapshot, notebook_audit),
        "practicalReferences": practical_spreadsheet_metadata(practical_spreadsheet_snapshot, practical_spreadsheet_audit),
        "practicalTemplateReferences": practical_template_metadata(practical_template_snapshot, practical_template_audit),
        "practicalRubricReferences": practical_rubric_metadata(practical_rubric_snapshot, practical_rubric_audit),
        "theory": theory,
        "officialTheoryMock": official_mock,
        "practical": practical,
        "officialPracticalMockCodes": ["1.1.1", "2.1.1", "2.2.1", "3.1.1", "3.2.1", "4.2.1"],
        "stats": {
            "theoryQuestions": len(theory),
            "officialTheoryAnswers": sum(q["answerSource"] == "official" for q in theory),
            "derivedTheoryAnswers": sum(q["answerSource"] == "derived" for q in theory),
            "practicalTasks": len(practical),
            "notebookTasks": sum(task["hasNotebook"] for task in practical),
            "officialPracticalAnswers": sum(task["answerSource"] == "official" for task in practical),
        },
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    write_json(args.theory_reference_audit, reference_audit)
    write_json(args.notebook_solution_audit, notebook_audit)
    print(json.dumps(payload["stats"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
