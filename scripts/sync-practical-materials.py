#!/usr/bin/env python3
"""Attach every official practical material to content.json and copy browser-safe downloads."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
from pathlib import Path
from urllib.parse import quote, unquote

from docx import Document
from openpyxl import load_workbook


SKIP_SUFFIXES = {".pyc"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def encoded_path(relative: Path) -> str:
    return "/".join(quote(part) for part in relative.parts)


def manifest_sha256(lines: list[str]) -> str:
    return hashlib.sha256("\n".join(lines).encode("utf-8")).hexdigest()


def workbook_note(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            sheets.append(f"{sheet.title}：{sheet.max_row} 行 × {sheet.max_column} 列")
        return "；".join(sheets)
    finally:
        workbook.close()


def csv_note(path: Path) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as stream:
                reader = csv.reader(stream)
                header = next(reader, [])
                rows = sum(1 for _ in reader)
            return f"{rows} 行数据 × {len(header)} 列；字段：{'、'.join(header[:5])}{'…' if len(header) > 5 else ''}"
        except UnicodeDecodeError:
            continue
    return "CSV 官方数据集"


def docx_note(path: Path) -> str:
    document = Document(path)
    paragraphs = sum(bool(paragraph.text.strip()) for paragraph in document.paragraphs)
    return f"官方配套文档；{paragraphs} 个正文段落，{len(document.tables)} 个表格"


def docx_preview(path: Path) -> list[dict[str, str]]:
    """Extract the official Word body for faithful, read-only browser display."""
    document = Document(path)
    preview = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "").lower()
        has_bold_text = any(run.bold and run.text.strip() for run in paragraph.runs)
        if style_name.startswith("heading") or has_bold_text:
            kind = "heading"
        elif "list" in style_name:
            kind = "list"
        else:
            kind = "paragraph"
        preview.append({"kind": kind, "text": text})
    return preview


def notebook_note(path: Path) -> str:
    notebook = json.loads(path.read_text(encoding="utf-8"))
    code_cells = sum(cell.get("cell_type") == "code" for cell in notebook.get("cells", []))
    return f"官方 Notebook 模板；{code_cells} 个代码单元"


def material_kind(path: Path) -> tuple[str, str]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        return "spreadsheet", "Excel 数据集"
    if suffix == ".csv":
        return "dataset", "CSV 数据集"
    if suffix in {".doc", ".docx"}:
        return "document", "配套文档"
    if suffix == ".ipynb":
        return "notebook", "Notebook 模板"
    if suffix == ".onnx":
        return "model", "ONNX 模型"
    if suffix in IMAGE_SUFFIXES:
        return "image", "测试图片"
    if suffix == ".py":
        return "code", "Python 代码"
    if suffix == ".txt":
        return "text", "标签文件"
    return "other", "配套文件"


def material_note(path: Path, kind: str) -> str:
    if kind == "spreadsheet":
        return workbook_note(path)
    if kind == "dataset":
        return csv_note(path)
    if kind == "document" and path.suffix.lower() == ".docx":
        return docx_note(path)
    if kind == "notebook":
        return notebook_note(path)
    if kind == "model":
        return "官方 ONNX 模型；进入实验室后按需加载浏览器兼容版本"
    return material_kind(path)[1]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path.cwd().parent / "3级评价指导手册（上网）" / "05-评价指导手册（上网）")
    parser.add_argument("--content", type=Path, default=Path("data/content.json"))
    parser.add_argument("--public", type=Path, default=Path("public/materials"))
    parser.add_argument("--audit", type=Path, default=Path("data/references/practical-material-audit.json"))
    args = parser.parse_args()

    content = json.loads(args.content.read_text(encoding="utf-8"))
    lab_manifest_path = Path("public/labs/manifest.json")
    lab_manifest_data = json.loads(lab_manifest_path.read_text(encoding="utf-8")) if lab_manifest_path.exists() else {"models": {}}
    browser_models = {model["originalSha256"]: model for model in lab_manifest_data.get("models", {}).values()}
    material_root = next((path for path in args.source.rglob("*") if path.is_dir() and path.name.endswith("_sucai")), None)
    if material_root is None:
        raise SystemExit("Cannot find official practical material directory")
    directories = {path.name.removesuffix("-素材"): path for path in material_root.iterdir() if path.is_dir() and path.name.endswith("-素材")}
    task_codes = {task["code"] for task in content["practical"]}
    orphan_directories = sorted(set(directories) - task_codes)
    tasks_without_source_directory = sorted(task_codes - set(directories))
    if orphan_directories:
        raise SystemExit(f"Official material directories are not mapped to tasks: {orphan_directories}")
    args.public.mkdir(parents=True, exist_ok=True)

    audit_tasks = []
    kind_counts: dict[str, int] = {}
    total_files = 0
    total_bytes = 0
    exact_browser_copies = 0
    verified_browser_deliveries = 0
    source_manifest: list[str] = []
    for task in content["practical"]:
        code = task["code"]
        source_dir = directories.get(code)
        entries = []
        if source_dir:
            for source in sorted(source_dir.rglob("*")):
                if not source.is_file() or source.suffix.lower() in SKIP_SUFFIXES or "__pycache__" in source.parts:
                    continue
                relative = source.relative_to(source_dir)
                kind, label = material_kind(source)
                url = None
                if kind == "notebook":
                    url = f"/lite/files/{code}/{encoded_path(relative)}"
                elif kind == "model":
                    url = None
                elif task["hasNotebook"] and kind not in {"document"}:
                    url = f"/labs/{code}/{encoded_path(relative)}"
                else:
                    target = args.public / code / relative
                    target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(source, target)
                    url = f"/materials/{code}/{encoded_path(relative)}"
                entry = {
                    "name": str(relative).replace("\\", "/"),
                    "kind": kind,
                    "label": label,
                    "size": source.stat().st_size,
                    "sha256": sha256(source),
                    "note": material_note(source, kind),
                    "url": url,
                }
                if task["group"] in {"4.1", "4.2"} and kind == "document" and source.suffix.lower() == ".docx":
                    entry["preview"] = docx_preview(source)
                entries.append(entry)
                source_manifest.append(f"{code}/{relative.as_posix()}\0{entry['size']}\0{entry['sha256']}")
                kind_counts[kind] = kind_counts.get(kind, 0) + 1
                total_files += 1
                total_bytes += entry["size"]
        task["materials"] = entries
        audited_entries = []
        for entry in entries:
            if entry["kind"] == "notebook":
                delivery = "jupyterlite-generated-template"
                published = args.public.parent / unquote(entry["url"]).lstrip("/")
                published_exists = published.is_file()
                published_exact = None
                delivery_verified = published_exists
            elif entry["kind"] == "model":
                delivery = "chunked-browser-cache"
                model = browser_models.get(entry["sha256"])
                chunks = model.get("chunks", []) if model else []
                delivery_verified = bool(chunks) and all(
                    (args.public.parent / chunk["url"].lstrip("/")).is_file()
                    and (args.public.parent / chunk["url"].lstrip("/")).stat().st_size == chunk["size"]
                    and sha256(args.public.parent / chunk["url"].lstrip("/")) == chunk["sha256"]
                    for chunk in chunks
                )
                published_exists = delivery_verified
                published_exact = None
            else:
                delivery = "lab-runtime" if entry["url"].startswith("/labs/") else "direct-download"
                published = args.public.parent / unquote(entry["url"]).lstrip("/")
                published_exists = published.is_file()
                published_exact = published_exists and published.stat().st_size == entry["size"] and sha256(published) == entry["sha256"]
                delivery_verified = published_exact
                exact_browser_copies += int(published_exact)
            verified_browser_deliveries += int(delivery_verified)
            audited_entries.append({
                **entry,
                "delivery": delivery,
                "publishedExists": published_exists,
                "publishedExactCopy": published_exact,
                "deliveryVerified": delivery_verified,
            })
        audit_tasks.append({
            "code": code,
            "sourceDirectory": source_dir.name if source_dir else None,
            "files": len(entries),
            "sourceStatus": "official-files-registered" if entries else "official-directory-not-provided",
            "allSourceFilesRegistered": len(entries) == sum(1 for path in source_dir.rglob("*") if path.is_file() and path.suffix.lower() not in SKIP_SUFFIXES and "__pycache__" not in path.parts) if source_dir else True,
            "materials": audited_entries,
        })

    content["version"] = "2026-07-19-v7"
    args.content.write_text(json.dumps(content, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    if lab_manifest_path.exists():
        manifest = json.loads(lab_manifest_path.read_text(encoding="utf-8"))
        manifest["version"] = content["version"]
        lab_manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    audit = {
        "contentVersion": content["version"],
        "materialRoot": material_root.name,
        "summary": {
            "tasks": len(content["practical"]),
            "sourceDirectories": len(directories),
            "tasksWithMaterials": sum(bool(item["files"]) for item in audit_tasks),
            "tasksWithoutOfficialFiles": tasks_without_source_directory,
            "orphanSourceDirectories": orphan_directories,
            "materialFiles": total_files,
            "materialBytes": total_bytes,
            "sourceTreeSha256": manifest_sha256(sorted(source_manifest)),
            "sourceFilesRegistered": sum(item["files"] for item in audit_tasks if item["allSourceFilesRegistered"]),
            "exactBrowserCopies": exact_browser_copies,
            "browserDeliveriesVerified": verified_browser_deliveries,
            "tasksFullyRegistered": sum(item["allSourceFilesRegistered"] for item in audit_tasks),
            "kinds": kind_counts,
        },
        "tasks": audit_tasks,
    }
    args.audit.parent.mkdir(parents=True, exist_ok=True)
    args.audit.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit["summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
