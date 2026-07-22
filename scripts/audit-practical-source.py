#!/usr/bin/env python3
"""Audit the 40 practical tasks against a DOCX copy of the official review manual."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document


def compact(value: str) -> str:
    return re.sub(r"[\s\x00-\x1f；。]", "", value).replace("/", "")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def official_rubric(table) -> list[dict]:
    result = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        if re.fullmatch(r"M\d+", cells[0]):
            result.append({"id": cells[0], "points": float(cells[1]), "description": cells[2]})
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manual-docx", type=Path, required=True)
    parser.add_argument("--manual-source", type=Path)
    parser.add_argument("--content", type=Path, default=Path("data/content.json"))
    parser.add_argument("--output", type=Path, default=Path("data/references/practical-source-audit.json"))
    args = parser.parse_args()

    content = json.loads(args.content.read_text(encoding="utf-8"))
    tasks = content["practical"]
    document = Document(args.manual_docx)
    if len(tasks) != 40 or len(document.tables) != 40:
        raise SystemExit(f"Expected 40 tasks and 40 scoring tables, got {len(tasks)} and {len(document.tables)}")

    paragraphs = [" ".join(paragraph.text.split()) for paragraph in document.paragraphs]
    markers = [index for index, text in enumerate(paragraphs) if text.startswith("试题名称")]
    if len(markers) != 80:
        raise SystemExit(f"Expected 80 task-name markers, got {len(markers)}")

    results = []
    for index, task in enumerate(tasks):
        marker = markers[index * 2]
        title = re.sub(r"^试题名称[:：]?", "", paragraphs[marker]).strip()
        work_marker = next((pos for pos in range(marker + 1, markers[index * 2 + 1]) if paragraphs[pos].startswith("2.") and "工作任务" in paragraphs[pos]), None)
        skill_marker = next((pos for pos in range((work_marker or marker) + 1, markers[index * 2 + 1]) if paragraphs[pos].startswith("3.") and "技能要求" in paragraphs[pos]), None)
        official_body = "\n".join(text for text in paragraphs[(work_marker or marker) + 1 : skill_marker] if text) if skill_marker else ""
        official_items = official_rubric(document.tables[index])
        current_items = task["rubric"]
        rubric_checks = []
        for official, current in zip(official_items, current_items):
            rubric_checks.append({
                "id": official["id"],
                "idMatch": official["id"] == current["id"],
                "pointsMatch": official["points"] == float(current["points"]),
                "descriptionMatch": compact(official["description"]) == compact(current["description"]),
            })
        results.append({
            "code": task["code"],
            "titleMatch": compact(title) == compact(task["title"]),
            "descriptionMatch": compact(official_body) == compact(task["description"]),
            "officialRubricItems": len(official_items),
            "siteRubricItems": len(current_items),
            "rubricCountMatch": len(official_items) == len(current_items),
            "rubricChecks": rubric_checks,
        })

    summary = {
        "tasks": len(results),
        "titleMatches": sum(item["titleMatch"] for item in results),
        "descriptionMatches": sum(item["descriptionMatch"] for item in results),
        "rubricCountMatches": sum(item["rubricCountMatch"] for item in results),
        "rubricItemChecks": sum(len(item["rubricChecks"]) for item in results),
        "rubricItemExactMatches": sum(
            check["idMatch"] and check["pointsMatch"] and check["descriptionMatch"]
            for item in results for check in item["rubricChecks"]
        ),
    }
    payload = {
        "manualSourceSha256": digest(args.manual_source) if args.manual_source else None,
        "convertedDocxSha256": digest(args.manual_docx),
        "summary": summary,
        "tasks": results,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
